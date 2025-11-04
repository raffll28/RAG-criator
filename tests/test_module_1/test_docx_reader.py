"""
Testes para DOCXReader
"""

import pytest
from pathlib import Path
import tempfile
import os
from docx import Document as DocxDoc
from docx.shared import Inches

from module_1_ingestion.readers import DOCXReader, Document


class TestDOCXReader:
    """Testes para a classe DOCXReader."""
    
    def setup_method(self):
        """Setup executado antes de cada teste."""
        self.reader = DOCXReader()
        self.temp_dir = tempfile.mkdtemp()
    
    def teardown_method(self):
        """Cleanup executado após cada teste."""
        import shutil
        if os.path.exists(self.temp_dir):
            shutil.rmtree(self.temp_dir)
    
    def _create_test_docx(self, filename: str, paragraphs: list, include_table: bool = False) -> Path:
        """
        Cria um DOCX de teste.
        
        Args:
            filename: Nome do arquivo
            paragraphs: Lista de parágrafos
            include_table: Se True, adiciona uma tabela de teste
            
        Returns:
            Path do arquivo criado
        """
        file_path = Path(self.temp_dir) / filename
        doc = DocxDoc()
        
        for para_text in paragraphs:
            doc.add_paragraph(para_text)
        
        if include_table:
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Header 1"
            table.cell(0, 1).text = "Header 2"
            table.cell(1, 0).text = "Value 1"
            table.cell(1, 1).text = "Value 2"
        
        doc.save(file_path)
        return file_path
    
    def test_reader_initialization(self):
        """Testa a inicialização do DOCXReader."""
        reader = DOCXReader()
        assert reader.include_tables is True
        assert ".docx" in reader.supported_extensions
    
    def test_supports_docx_files(self):
        """Testa se o reader identifica arquivos DOCX."""
        assert self.reader.supports(Path("test.docx"))
        assert not self.reader.supports(Path("test.txt"))
        assert not self.reader.supports(Path("test.pdf"))
    
    def test_read_simple_docx(self):
        """Testa leitura de DOCX simples."""
        paragraphs = ["Hello World", "This is a test document", "Third paragraph"]
        test_file = self._create_test_docx("test.docx", paragraphs)
        
        document = self.reader.read(test_file)
        
        assert isinstance(document, Document)
        assert "Hello World" in document.content
        assert "test document" in document.content
        assert document.metadata["file_extension"] == ".docx"
        assert document.metadata["paragraphs_count"] == 3
        assert document.metadata["reader_type"] == "DOCXReader"
    
    def test_read_docx_with_table(self):
        """Testa leitura de DOCX com tabela."""
        paragraphs = ["Document with table"]
        test_file = self._create_test_docx("table.docx", paragraphs, include_table=True)
        
        document = self.reader.read(test_file)
        
        assert document.metadata["tables_count"] == 1
        assert "Header 1" in document.content
        assert "Value 1" in document.content
    
    def test_read_nonexistent_docx(self):
        """Testa leitura de DOCX que não existe."""
        test_file = Path(self.temp_dir) / "nonexistent.docx"
        
        with pytest.raises(FileNotFoundError):
            self.reader.read(test_file)
    
    def test_docx_metadata_extraction(self):
        """Testa extração de metadados do DOCX."""
        paragraphs = ["Test paragraph"]
        test_file = self._create_test_docx("metadata.docx", paragraphs)
        
        document = self.reader.read(test_file)
        metadata = document.metadata
        
        # Verifica campos obrigatórios
        required_fields = [
            "paragraphs_count", "tables_count", "sections_count",
            "word_count", "content_length", "content_hash"
        ]
        
        for field in required_fields:
            assert field in metadata, f"Missing metadata field: {field}"
        
        assert metadata["paragraphs_count"] > 0
    
    def test_empty_docx(self):
        """Testa leitura de DOCX vazio."""
        test_file = self._create_test_docx("empty.docx", [])
        
        document = self.reader.read(test_file)
        
        assert document.allow_empty is True
        assert document.metadata["is_empty"] is True
        assert document.metadata["paragraphs_count"] == 0
    
    def test_docx_without_tables(self):
        """Testa reader configurado para não extrair tabelas."""
        reader = DOCXReader(include_tables=False)
        paragraphs = ["Content"]
        test_file = self._create_test_docx("notables.docx", paragraphs, include_table=True)
        
        document = reader.read(test_file)
        
        # Tabela existe mas não deve ser extraída
        assert document.metadata["tables_count"] == 0

